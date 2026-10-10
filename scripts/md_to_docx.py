"""Convert the career CV markdown to .docx (compact, ATS-friendly).

Usage: python scripts/md_to_docx.py <input.md> <output.docx>

Handles the subset used by docs/career/*.md: # title, ## sections, **bold**
inline, "- " bullets, "---" rules (skipped). Not a general markdown converter.
"""

from __future__ import annotations

import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** spans."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1


def convert(src: str, dst: str) -> None:
    doc = Document()

    # Compact margins + base font.
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)

    with open(src, encoding="utf-8") as handle:
        lines = handle.read().splitlines()

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, stripped[2:])
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(18)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            add_runs(p, stripped[3:].upper())
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x1A, 0x2A, 0x4A)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)

    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: python scripts/md_to_docx.py <in.md> <out.docx>")
        raise SystemExit(2)
    convert(sys.argv[1], sys.argv[2])
